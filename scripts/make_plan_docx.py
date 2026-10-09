# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from datetime import date

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)


def set_run_font(run, size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else (14 if level == 2 else 12), bold=True)
    return p


def add_para(text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_check(text, done=True):
    p = doc.add_paragraph()
    mark = "☑" if done else "☐"
    run = p.add_run(f"{mark}  {text}")
    set_run_font(run, size=11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.line_spacing = 1.2
    return p


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(title.add_run("会展 3D 大屏 · FastAPI · AI 对话操控"), size=18, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(sub.add_run("两个月学习与上线实战计划（修订版）"), size=14, bold=True)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(
    meta.add_run(
        f"项目：convention-3d + api-convention　　修订日期：{date.today().isoformat()}　　可打印"
    ),
    size=10,
    color=(80, 80, 80),
)

add_heading_cn("一、项目目标与边界", 1)
add_para(
    "目标：在现有 Vue3 + Three.js 会展大屏上，用自己写的 Python FastAPI 提供数据；"
    "再用中文对话查询数据、操控 3D 场景；最终部署到公网，获得可写进简历的上线经验。"
)
add_para(
    "边界：不重写前端 3D、不训练大模型、第一阶段不用 LangChain、不接真实 ICC 视频平台。"
    "AI 只通过白名单指令调用已有 eventHub。"
)

add_heading_cn("二、当前仓库现状（对照代码）", 1)
add_para(
    "前端 convention-3d：综合态势 / 安防态势 / 设备运行 / 能源管理 / 会展信息；"
    "ControlPanel 已通过 mitt 控制天气、昼夜、场馆、分层、热力图、漫游。"
)
add_para(
    "后端 api-convention：FastAPI + SQLAlchemy + SQLite 种子数据；接口包括 halls、devices、"
    "exhibitions、alarms、work-orders、energy、parking、crowd、heatmap、emergency。"
)
add_para(
    "联调：Vite 代理 /api → 127.0.0.1:8000；前端 http.js + convention.js + screenStats.js "
    "已把五页图表/列表接到后端。"
)

add_heading_cn("三、已完成内容（打勾）", 1)

add_heading_cn("3.1 第 1 周：后端地基与代理", 2)
for t in [
    "搭建 api-convention（FastAPI）与本地启动脚本 start.bat",
    "Python 3.12 虚拟环境 .venv，依赖写入 requirements.txt",
    "GET /health 健康检查可用",
    "CORS 允许 localhost:3000 / 127.0.0.1:3000",
    "Vite 代理配置到 http://127.0.0.1:8000，保留 /api 前缀",
    "新建前端 src/api/http.js（axios，适配 Vite）",
]:
    add_check(t, True)

add_heading_cn("3.2 第 2 周：数据库与业务接口", 2)
for t in [
    "SQLAlchemy 模型：Hall / Device / Exhibition / ExhibitionEvent / Alarm / WorkOrder / EnergyReading / ParkingStat / CrowdStat / HeatmapPoint / EmergencyRecord",
    "seed.py 种子数据（场馆、设备、展会、告警、工单、能耗、车位、人流、热力图、应急）",
    "GET /api/halls、/api/halls/{name}",
    "GET /api/devices、/api/devices/count",
    "GET /api/exhibitions、/api/exhibition-events",
    "GET /api/alarms（支持等级/年月过滤）",
    "GET /api/work-orders",
    "GET /api/energy",
    "GET /api/parking、/api/crowd、/api/heatmap、/api/emergency",
    "修复 /api/emergency 返回列表（原集合无法 JSON 序列化）",
]:
    add_check(t, True)

add_heading_cn("3.3 第 3 周：五页前端对接后端数据", 2)
for t in [
    "src/api/convention.js 统一封装接口调用与缓存",
    "src/utils/screenStats.js 汇总统计与图表数据",
    "综合态势页：展馆总览、会展/车位/人流、告警、能耗、工单",
    "设备运行页：设备统计、告警、工单类型/列表/超时 TOP",
    "能源管理页：用电/用水、占比、能耗告警与工单、分析与指标",
    "安防态势页：安防设备与告警、通行车位与人流、应急概况",
    "会展信息页：展会统计、排期与会议安排、服务信息",
    "顶栏未处理告警数接入 /api/alarms",
    "浏览器实测：页面数字与列表来自后端接口",
]:
    add_check(t, True)

add_heading_cn("3.4 工程与启动问题修复", 2)
for t in [
    "修复 start.bat：去掉失效的 Python3.14 硬编码路径，自动找本机 Python",
    "重建损坏的 .venv（指向已删除解释器的问题）",
    "错误时 pause，避免双击窗口一闪而过",
]:
    add_check(t, True)

add_heading_cn("四、未开始内容（待办清单）", 1)
add_para(
    "以下按阶段推进。建议工作日每天 3 小时，周末 4～5 小时。每天收工以可运行结果为准。"
)

add_heading_cn("4.1 阶段 A：鉴权与接口规范（约 4～5 天）", 2)
for t in [
    "JWT 登录接口 POST /api/auth/login（演示账号 admin）",
    "密码哈希存储；Bearer Token 保护业务接口",
    "前端登录页或启动时写入 token，请求头自动带 Authorization",
    "统一响应格式约定（可选：{code,msg,data}）与 401 处理",
    "README 写清：如何启动前后端、演示账号、常用接口",
]:
    add_check(t, False)

add_heading_cn("4.2 阶段 B：场景指令通道——先不接 AI（约 7 天）", 2)
add_para("目标：人和接口共用同一套 command，为后面大模型 Tool Calling 打地基。")
for t in [
    "前端写 aiCommandBus：JSON action 映射到 eventHub / router",
    "白名单：navigate / weather / dayNight / building / floor / heatmap / roam / openDevice",
    "后端 POST /api/scene/command 回声或校验后返回指令",
    "非法指令拒绝；GET /api/scene/state 回传当前场景状态",
    "ControlPanel 点击也走同一套 command 格式",
    "手测 15 条指令（下雨、夜晚、1号馆、分层、热力图、漫游、切页）",
    "不点面板，只靠接口把场景玩一遍（周演示）",
]:
    add_check(t, False)

add_heading_cn("4.3 阶段 C：接入大模型对话（约 7 天）", 2)
add_para("推荐 DeepSeek（OpenAI 兼容接口）。禁止训练模型；禁止让模型生成 Three.js 代码。")
for t in [
    "申请 API Key，key 只放 .env，永不进 Git",
    "scripts/chat_hello.py 终端对话跑通",
    "Function Calling：query_alarms / query_energy 等查询工具接数据库",
    "把场景指令注册成 tools（改天气、切馆、切页等）",
    "POST /api/chat 非流式对话；系统提示词限制只能用给定工具",
    "前端 AiChatPanel：消息历史、loading、错误提示",
    "验收话术：「改成雨天」真下雨；「今天有哪些告警」返回库里的数",
]:
    add_check(t, False)

add_heading_cn("4.4 阶段 D：对话产品化（约 7 天）", 2)
for t in [
    "SSE 流式输出 /api/chat/stream，前端逐字显示",
    "多工具串联：去1号馆 + 开热力图 + 查告警",
    "chat_sessions / chat_messages 落库，刷新不丢会话",
    "提示词打磨：禁止编造数字、禁止白名单外操作",
    "快捷问题 chips：看看告警 / 下雨 / 漫游路线一",
    "对话日志：原话、工具调用、耗时；修 10 个 bad case",
    "找朋友试用并录 2 分钟演示视频",
]:
    add_check(t, False)

add_heading_cn("4.5 阶段 E：上线准备与公网实战（约 14 天）", 2)
for t in [
    "冻死 requirements.txt；写 Dockerfile，本机 docker run 起 API",
    "vite build + Nginx：静态资源 + /api 反代",
    "轻量服务器（2核2G 即可）+ 安全组 80/443",
    "域名解析 + HTTPS（Caddy 或 Nginx + 证书）",
    "生产环境导入种子数据；CORS 改为正式域名",
    "LLM key 只放服务器；本地与生产环境分离",
    "公网冒烟：登录、五页、3D、对话各测一遍",
    "发布 v1.0，发给至少 3 人真用",
    "聊天限流，防止 token 被刷",
    "访问日志与错误日志，用日志排除一次故障",
    "README：启动、环境变量、演示账号、AI 能力边界",
    "可选加分一项：告警一句话解读 或 会展简介检索（二选一）",
    "5 分钟演示稿 + 录屏；写简历项目描述",
]:
    add_check(t, False)

add_heading_cn("4.6 阶段 F：收尾结案（约 5 天）", 2)
for t in [
    "pytest：登录、设备列表、指令白名单、chat tool（5～8 个）",
    "清理死代码（event copy、dialog copy 等）",
    "统计一个月服务器 + DeepSeek 费用",
    "复盘文档：做成了什么 / 没做成什么 / 下一步",
    "冻结版本，GitHub 脱敏后公开（可选）",
]:
    add_check(t, False)

add_heading_cn("五、验收标准（全部完成时）", 1)
for t in [
    "公网 HTTPS 可打开现有 3D 会展页",
    "五个导航页数据来自 FastAPI，不再写死",
    "能登录，演示账号写在 README",
    "对话能切页、改天气/昼夜/馆/楼层/热力图/漫游",
    "对话能查告警、能耗、会展，数字与接口一致",
]:
    add_check(t, False)

add_heading_cn("六、建议时间节奏（从今天起重新排）", 1)
add_para("已完成：第 1～3 周核心（后端接口 + 五页数据对接）。")
add_para(
    "接下来：约 1～2 周做阶段 A+B（鉴权 + 场景指令）；约 3～4 周做阶段 C+D（大模型对话）；"
    "约 5～6 周做阶段 E（上线）；最后几天做阶段 F（收尾）。"
)
add_para(
    "卡壳铁律：同一问题连续卡 2 小时就降级（先 mock、先非流式、先 SQLite）；"
    "周日必须能给自己演示；周目标不能丢。"
)

add_heading_cn("七、每日自检栏（打印后可手写）", 1)
add_para("日期：__________　　阶段：A / B / C / D / E / F　　投入小时：____")
add_para("今天完成的条目（抄序号或简述）：")
add_para("______________________________________________________________")
add_para("______________________________________________________________")
add_para("可运行证据（curl /docs / 页面截图 / 录屏）：________________")
add_para("卡住的问题与下一步：________________________________________")

footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(
    footer.add_run("— 文档结束 · 对照仓库 convention-3d + api-convention 修订 —"),
    size=9,
    color=(120, 120, 120),
)

out = r"c:\Users\luyin\Desktop\lily\convention\会展3D-FastAPI-AI两个月计划-修订版.docx"
doc.save(out)
print(out)
